#!/usr/bin/env python3
"""
Convert one or more flat JSONL files (one JSON object per line, no nested
objects/arrays - like questions_formation/tache2.jsonl) into a single Excel
workbook, one sheet per input file.

    pip install pandas openpyxl

    python3 jsonl_to_xlsx.py questions_formation/tache2.jsonl
    python3 jsonl_to_xlsx.py questions_formation/tache2.jsonl questions_formation/tache3.jsonl -o questions_formation.xlsx
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

import pandas as pd
from openpyxl.utils import get_column_letter


def load_jsonl(path: Path) -> list[dict]:
    rows: list[dict] = []
    with path.open(encoding="utf-8") as fh:
        for lineno, line in enumerate(fh, 1):
            line = line.strip()
            if not line:
                continue
            try:
                rows.append(json.loads(line))
            except json.JSONDecodeError as exc:
                sys.exit(f"{path}:{lineno}: invalid JSON - {exc}")
    return rows


def autosize_columns(ws, df: pd.DataFrame, max_width: int = 60) -> None:
    """Excel doesn't autosize columns on its own; widen each to fit its
    longest cell, capped since free-text columns can run to paragraphs."""
    for i, col in enumerate(df.columns, start=1):
        longest = max([len(str(col))] + [len(str(v)) for v in df[col]])
        ws.column_dimensions[get_column_letter(i)].width = min(longest + 2, max_width)


def force_text_columns(ws, df: pd.DataFrame, cols: set[str]) -> None:
    """Mark these columns as text in Excel so it never reinterprets them.

    The scrapers' ids look numeric but carry a significant leading zero
    ("0320260820401"). Without an explicit text format, editing the cell in
    Excel can silently renumber it to 320260820401 and destroy the 2-digit
    source prefix.
    """
    for i, col in enumerate(df.columns, start=1):
        if col not in cols:
            continue
        letter = get_column_letter(i)
        for row in range(2, ws.max_row + 1):        # skip the header
            ws[f"{letter}{row}"].number_format = "@"


def convert(paths: list[Path], out_path: Path, text_cols: set[str]) -> None:
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        for path in paths:
            rows = load_jsonl(path)
            if not rows:
                print(f"  ! {path}: no rows, skipping", file=sys.stderr)
                continue
            df = pd.DataFrame(rows)
            sheet_name = path.stem[:31]                # Excel's sheet-name length limit
            df.to_excel(writer, sheet_name=sheet_name, index=False)
            autosize_columns(writer.sheets[sheet_name], df)
            force_text_columns(writer.sheets[sheet_name], df, text_cols)
            print(f"{path} -> sheet {sheet_name!r} ({len(df)} rows, {len(df.columns)} cols)")

    print(f"\nwrote {out_path}")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("jsonl", nargs="+", type=Path, help="one or more flat .jsonl files")
    ap.add_argument("-o", "--out", type=Path,
                    help="output .xlsx path (default: <first file>.xlsx)")
    ap.add_argument("--text-cols", default="id",
                    help="comma-separated columns to format as text in Excel, "
                         "protecting significant leading zeros (default: id)")
    args = ap.parse_args()

    for p in args.jsonl:
        if not p.exists():
            sys.exit(f"not found: {p}")

    out_path = args.out or args.jsonl[0].with_suffix(".xlsx")
    text_cols = {c.strip() for c in args.text_cols.split(",") if c.strip()}
    convert(args.jsonl, out_path, text_cols)


if __name__ == "__main__":
    main()
