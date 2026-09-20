"""Admin: vocabulary management and the labelling workspace.

Everything here is guarded by `current_admin`, which 403s a signed-in
non-admin.

What this replaces
------------------
    load_vocabulary.py --prune   add / delete labels, with the in-use check
    the hand-written repoint     merging one label into another
    export -> edit -> load_labels.py   for correcting a label by hand

The CLI tools stay useful for bulk work driven by llm.py; this is for the
one-at-a-time decisions that a terminal makes tedious.
"""
from __future__ import annotations

from typing import Optional

import io

from fastapi import APIRouter, Depends, HTTPException, Query, Response, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from ..auth import current_admin
from ..db import get_db
from ..models import CoreSubject, Fingerprint, Theme, User

router = APIRouter(prefix="/api/admin", tags=["admin"],
                   dependencies=[Depends(current_admin)])


def _norm(name: str) -> str:
    """Trim and collapse whitespace. Stops ' lessons' and 'lessons ' becoming
    two labels that look identical in the UI."""
    return " ".join(name.split())


def _key(name: str) -> str:
    """Case- and apostrophe-insensitive comparison form, matching
    load_labels.py - so a duplicate check here agrees with what a load would
    actually collide with."""
    return _norm(name).replace("’", "'").casefold()


# ----------------------------------------------------------------- vocabulary


class ThemeIn(BaseModel):
    tache: int = Field(ge=2, le=3)
    name: str = Field(min_length=1, max_length=40)


class RenameIn(BaseModel):
    name: str = Field(min_length=1, max_length=60)


class CoreSubjectIn(BaseModel):
    theme_id: int
    name: str = Field(min_length=1, max_length=60)


class MergeIn(BaseModel):
    into_id: int


@router.get("/vocabulary")
def vocabulary(tache: int = Query(2, ge=2, le=3),
               db: Session = Depends(get_db)) -> dict:
    """Themes and their core subjects, with how many questions use each.

    The counts are the point: they are what tells you a label is dead wood, a
    near-duplicate of its neighbour, or the only thing holding a theme up.
    """
    used = dict(db.execute(
        select(Fingerprint.core_subject_id, func.count())
        .where(Fingerprint.core_subject_id.isnot(None))
        .group_by(Fingerprint.core_subject_id)).all())
    theme_used = dict(db.execute(
        select(Fingerprint.theme_id, func.count())
        .where(Fingerprint.theme_id.isnot(None))
        .group_by(Fingerprint.theme_id)).all())

    themes = db.scalars(select(Theme).where(Theme.tache == tache)
                        .order_by(Theme.name)).all()
    subjects: dict[int, list] = {}
    for cs in db.scalars(select(CoreSubject)
                         .where(CoreSubject.theme_id.in_([t.id for t in themes]))
                         .order_by(CoreSubject.name)).all():
        subjects.setdefault(cs.theme_id, []).append(
            {"id": cs.id, "name": cs.name, "usage": used.get(cs.id, 0)})

    return {
        "tache": tache,
        "themes": [{
            "id": t.id,
            "name": t.name,
            "usage": theme_used.get(t.id, 0),
            "core_subjects": subjects.get(t.id, []),
        } for t in themes],
    }


@router.post("/themes", status_code=status.HTTP_201_CREATED)
def create_theme(payload: ThemeIn, db: Session = Depends(get_db)) -> dict:
    name = _norm(payload.name)
    existing = db.scalars(select(Theme).where(Theme.tache == payload.tache)).all()
    if any(_key(t.name) == _key(name) for t in existing):
        raise HTTPException(409, f"tache {payload.tache} already has a theme "
                                 f"called {name!r}")
    t = Theme(tache=payload.tache, name=name)
    db.add(t)
    db.commit()
    return {"id": t.id, "name": t.name, "usage": 0, "core_subjects": []}


@router.patch("/themes/{theme_id}")
def rename_theme(theme_id: int, payload: RenameIn,
                 db: Session = Depends(get_db)) -> dict:
    t = db.get(Theme, theme_id)
    if t is None:
        raise HTTPException(404, "theme not found")
    name = _norm(payload.name)[:40]
    clash = db.scalars(select(Theme).where(Theme.tache == t.tache,
                                           Theme.id != t.id)).all()
    if any(_key(o.name) == _key(name) for o in clash):
        raise HTTPException(409, f"another theme is already called {name!r}")
    t.name = name
    db.commit()
    return {"id": t.id, "name": t.name}


@router.delete("/themes/{theme_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_theme(theme_id: int, db: Session = Depends(get_db)) -> Response:
    t = db.get(Theme, theme_id)
    if t is None:
        raise HTTPException(404, "theme not found")
    # refuse rather than cascade: deleting a theme with questions under it would
    # silently strip their labels, and there is no undo here
    n_q = db.scalar(select(func.count()).select_from(Fingerprint)
                    .where(Fingerprint.theme_id == theme_id)) or 0
    n_cs = db.scalar(select(func.count()).select_from(CoreSubject)
                     .where(CoreSubject.theme_id == theme_id)) or 0
    if n_q or n_cs:
        raise HTTPException(409, f"in use: {n_q} question(s) and {n_cs} core "
                                 f"subject(s). Move or delete those first")
    db.delete(t)
    db.commit()
    return Response(status_code=status.HTTP_204_NO_CONTENT)


@router.post("/core-subjects", status_code=status.HTTP_201_CREATED)
def create_core_subject(payload: CoreSubjectIn,
                        db: Session = Depends(get_db)) -> dict:
    theme = db.get(Theme, payload.theme_id)
    if theme is None:
        raise HTTPException(404, "theme not found")
    name = _norm(payload.name)
    siblings = db.scalars(select(CoreSubject)
                          .where(CoreSubject.theme_id == payload.theme_id)).all()
    if any(_key(c.name) == _key(name) for c in siblings):
        raise HTTPException(409, f"{theme.name} already has {name!r}")
    cs = CoreSubject(theme_id=payload.theme_id, name=name)
    db.add(cs)
    db.commit()
    return {"id": cs.id, "name": cs.name, "usage": 0}


@router.patch("/core-subjects/{cs_id}")
def update_core_subject(cs_id: int, payload: CoreSubjectIn,
                        db: Session = Depends(get_db)) -> dict:
    cs = db.get(CoreSubject, cs_id)
    if cs is None:
        raise HTTPException(404, "core subject not found")
    if db.get(Theme, payload.theme_id) is None:
        raise HTTPException(404, "theme not found")
    name = _norm(payload.name)
    siblings = db.scalars(select(CoreSubject)
                          .where(CoreSubject.theme_id == payload.theme_id,
                                 CoreSubject.id != cs_id)).all()
    if any(_key(c.name) == _key(name) for c in siblings):
        raise HTTPException(409, f"that theme already has {name!r}")

    moved = cs.theme_id != payload.theme_id
    cs.name, cs.theme_id = name, payload.theme_id
    if moved:
        # a question's theme and its core subject's theme must agree, or the
        # card would show a label belonging to a different theme
        db.execute(Fingerprint.__table__.update()
                   .where(Fingerprint.core_subject_id == cs_id)
                   .values(theme_id=payload.theme_id))
    db.commit()
    return {"id": cs.id, "name": cs.name, "theme_id": cs.theme_id, "moved": moved}


@router.delete("/core-subjects/{cs_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_core_subject(cs_id: int, db: Session = Depends(get_db)) -> Response:
    cs = db.get(CoreSubject, cs_id)
    if cs is None:
        raise HTTPException(404, "core subject not found")
    n = db.scalar(select(func.count()).select_from(Fingerprint)
                  .where(Fingerprint.core_subject_id == cs_id)) or 0
    if n:
        raise HTTPException(409, f"in use by {n} question(s). Merge it into "
                                 f"another label instead of deleting")
    db.delete(cs)
    db.commit()
    return Response(status_code=status.HTTP_204_NO_CONTENT)


@router.post("/core-subjects/{cs_id}/merge")
def merge_core_subject(cs_id: int, payload: MergeIn,
                       db: Session = Depends(get_db)) -> dict:
    """Repoint every question from one label to another, then delete the first.

    The move that `delete` refuses to do implicitly. Both labels must sit under
    the same theme - merging across themes would change what those questions
    are about, not just what they are called.
    """
    src, dst = db.get(CoreSubject, cs_id), db.get(CoreSubject, payload.into_id)
    if src is None or dst is None:
        raise HTTPException(404, "core subject not found")
    if src.id == dst.id:
        raise HTTPException(400, "cannot merge a label into itself")
    if src.theme_id != dst.theme_id:
        raise HTTPException(409, "both labels must be under the same theme")

    moved = db.execute(Fingerprint.__table__.update()
                       .where(Fingerprint.core_subject_id == src.id)
                       .values(core_subject_id=dst.id)).rowcount
    name = src.name
    db.delete(src)
    db.commit()
    return {"merged": name, "into": dst.name, "questions_moved": moved}


# ------------------------------------------------------------------ labelling


class LabelIn(BaseModel):
    # None clears the label; omitted leaves it alone. The two are different, so
    # the endpoint reads the raw body rather than trusting Pydantic defaults.
    theme_id: Optional[int] = None
    core_subject_id: Optional[int] = None
    abstract: Optional[str] = None


class BulkLabelIn(BaseModel):
    f_ids: list[int]
    theme_id: Optional[int] = None
    core_subject_id: Optional[int] = None


def _question_row(f: Fingerprint, theme: Optional[Theme],
                  cs: Optional[CoreSubject]) -> dict:
    return {
        "f_id": f.id, "text": f.text, "tache": f.tache,
        "abstract": f.abstract,
        "theme_id": f.theme_id, "theme": theme.name if theme else None,
        "core_subject_id": f.core_subject_id,
        "core_subject": cs.name if cs else None,
        "months_seen": f.months_seen, "total_sightings": f.total_sightings,
        "last_seen": f.last_seen,
    }


def _question_where(tache, q, theme_id, core_subject_id, unlabelled, inconsistent):
    """The filter behind both the page and its "select all matching" list.

    Shared so the two cannot drift: a button that claims to select what the
    page is showing has to mean the same thing by "showing".
    """
    where = [Fingerprint.tache == tache]
    if q:
        where.append(or_(Fingerprint.text.ilike(f"%{q}%"),
                         Fingerprint.abstract.ilike(f"%{q}%")))
    if theme_id:
        where.append(Fingerprint.theme_id == theme_id)
    if core_subject_id:
        where.append(Fingerprint.core_subject_id == core_subject_id)
    if inconsistent:
        # a pair no current endpoint can create, but older loads could: the
        # question says one theme and its core subject lives under another
        where.append(Fingerprint.core_subject_id.isnot(None))
        where.append(Fingerprint.theme_id != CoreSubject.theme_id)
    if unlabelled == "theme":
        where.append(Fingerprint.theme_id.is_(None))
    elif unlabelled == "core_subject":
        where.append(Fingerprint.core_subject_id.is_(None))
    elif unlabelled == "abstract":
        where.append(Fingerprint.abstract.is_(None))
    return where


@router.get("/questions/ids")
def question_ids(
    tache: int = Query(2, ge=2, le=3),
    q: Optional[str] = None,
    theme_id: Optional[int] = None,
    core_subject_id: Optional[int] = None,
    unlabelled: Optional[str] = Query(None, pattern="^(theme|abstract|core_subject)$"),
    inconsistent: bool = Query(False),
    db: Session = Depends(get_db),
) -> dict:
    """Every question id matching the filter, for "select all matching".

    Capped at the same ceiling a job is: selecting 5,000 questions and then
    being told the run will only cover 600 would be a worse experience than
    being told now. `total` is the unclipped count, so the UI can say so.
    """
    from ..llm_runner import MAX_ROWS

    where = _question_where(tache, q, theme_id, core_subject_id, unlabelled,
                            inconsistent)
    base = (select(Fingerprint.id)
            .outerjoin(CoreSubject, CoreSubject.id == Fingerprint.core_subject_id)
            .where(*where))
    total = db.scalar(select(func.count()).select_from(Fingerprint)
                      .outerjoin(CoreSubject,
                                 CoreSubject.id == Fingerprint.core_subject_id)
                      .where(*where)) or 0
    ids = [i for (i,) in db.execute(
        base.order_by(Fingerprint.months_seen.desc(), Fingerprint.id.desc())
            .limit(MAX_ROWS)).all()]
    return {"ids": ids, "total": total, "cap": MAX_ROWS, "capped": total > len(ids)}


EXPORT_COLUMNS = [
    ("id", 8), ("tache", 6), ("theme", 24), ("core_subject", 24),
    ("abstract", 34), ("text", 90), ("months_seen", 12),
    ("total_sightings", 14), ("first_seen", 11), ("last_seen", 11),
    ("fingerprint", 18),
]


@router.get("/questions/export")
def export_questions(
    tache: int = Query(2, ge=2, le=3),
    q: Optional[str] = None,
    theme_id: Optional[int] = None,
    core_subject_id: Optional[int] = None,
    unlabelled: Optional[str] = Query(None, pattern="^(theme|abstract|core_subject)$"),
    inconsistent: bool = Query(False),
    db: Session = Depends(get_db),
) -> StreamingResponse:
    """Every field of every question the filter matches, as a spreadsheet.

    Not capped and not paginated: an export is free and a partial one is a
    trap. Built server-side because the data is already here - shipping JSON to
    the browser and converting it there would mean a second definition of what
    a question's fields are.

    The labels are the resolved names, not the foreign keys, because the point
    of the file is to be read.
    """
    where = _question_where(tache, q, theme_id, core_subject_id, unlabelled,
                            inconsistent)
    rows = db.execute(
        select(Fingerprint, Theme.name, CoreSubject.name)
        .outerjoin(Theme, Theme.id == Fingerprint.theme_id)
        .outerjoin(CoreSubject, CoreSubject.id == Fingerprint.core_subject_id)
        .where(*where)
        .order_by(Fingerprint.months_seen.desc(), Fingerprint.id.desc())).all()

    wb = Workbook()
    ws = wb.active
    ws.title = f"Task {tache}"

    ws.append([name for name, _ in EXPORT_COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="top")
    # the header stays put while scrolling a few thousand rows, and a filter
    # row makes the file usable without anyone writing a formula
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(EXPORT_COLUMNS))}1"

    for f, theme, core_subject in rows:
        ws.append([
            f.id, f.tache, theme, core_subject, f.abstract, f.text,
            f.months_seen, f.total_sightings, f.first_seen, f.last_seen,
            # the hash is 64 characters of noise in a spreadsheet; the first 12
            # is enough to match rows against another export
            (f.fingerprint or "")[:12],
        ])

    for i, (_name, width) in enumerate(EXPORT_COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    # the question text is a paragraph; without wrapping every row is one line
    # running off the screen
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=6):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    bits = [str(len(rows)), f"tache{tache}"]
    if unlabelled:
        bits.append(f"no_{unlabelled}")
    if inconsistent:
        bits.append("mismatched")
    name = f"questions_{'_'.join(bits)}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{name}"'})


@router.get("/questions")
def list_questions(
    tache: int = Query(2, ge=2, le=3),
    q: Optional[str] = None,
    theme_id: Optional[int] = None,
    core_subject_id: Optional[int] = None,
    unlabelled: Optional[str] = Query(None, pattern="^(theme|abstract|core_subject)$"),
    inconsistent: bool = Query(False,
                               description="only questions whose core subject "
                                           "belongs to a different theme"),
    page: int = Query(1, ge=1),
    per_page: int = Query(50, ge=1, le=200),
    db: Session = Depends(get_db),
) -> dict:
    """Fingerprints with their labels - one row per question, not per month."""
    where = _question_where(tache, q, theme_id, core_subject_id, unlabelled,
                            inconsistent)

    # counted through the same joins as the rows. A filter that mentions
    # core_subjects without joining it turns the count into a cartesian product
    # - 70,522 instead of 4, with the right rows underneath it.
    joined = (select(func.count()).select_from(Fingerprint)
              .outerjoin(Theme, Theme.id == Fingerprint.theme_id)
              .outerjoin(CoreSubject, CoreSubject.id == Fingerprint.core_subject_id))
    total = db.scalar(joined.where(*where)) or 0
    rows = db.execute(
        select(Fingerprint, Theme, CoreSubject)
        .outerjoin(Theme, Theme.id == Fingerprint.theme_id)
        .outerjoin(CoreSubject, CoreSubject.id == Fingerprint.core_subject_id)
        .where(*where)
        .order_by(Fingerprint.months_seen.desc(), Fingerprint.id.desc())
        .offset((page - 1) * per_page).limit(per_page)).all()

    return {
        "items": [_question_row(f, t, c) for f, t, c in rows],
        "total": total, "page": page, "per_page": per_page,
    }


def _check_pair(db: Session, f: Fingerprint, theme_id: Optional[int],
                cs_id: Optional[int]) -> None:
    """A core subject must belong to the theme the question is being given."""
    if cs_id is None:
        return
    cs = db.get(CoreSubject, cs_id)
    if cs is None:
        raise HTTPException(404, "core subject not found")
    effective_theme = theme_id if theme_id is not None else f.theme_id
    if cs.theme_id != effective_theme:
        raise HTTPException(409, "that core subject belongs to a different theme")


@router.patch("/questions/{f_id}")
def set_labels(f_id: int, payload: dict, db: Session = Depends(get_db)) -> dict:
    """Set or clear labels on one question.

    Takes a raw dict so that an omitted key and an explicit null mean different
    things: leave alone, versus clear.
    """
    f = db.get(Fingerprint, f_id)
    if f is None:
        raise HTTPException(404, "question not found")

    theme_id = payload.get("theme_id", f.theme_id)
    cs_id = payload.get("core_subject_id", f.core_subject_id)
    if theme_id is not None and db.get(Theme, theme_id) is None:
        raise HTTPException(404, "theme not found")
    # changing the theme orphans a core subject that belonged to the old one
    if "theme_id" in payload and theme_id != f.theme_id and "core_subject_id" not in payload:
        cs_id = None
    _check_pair(db, f, theme_id, cs_id)

    f.theme_id, f.core_subject_id = theme_id, cs_id
    if "abstract" in payload:
        a = payload["abstract"]
        f.abstract = (a.strip()[:200] or None) if isinstance(a, str) else None
    db.commit()
    return _question_row(f, db.get(Theme, f.theme_id) if f.theme_id else None,
                         db.get(CoreSubject, f.core_subject_id) if f.core_subject_id else None)


@router.post("/questions/bulk")
def bulk_label(payload: BulkLabelIn, db: Session = Depends(get_db)) -> dict:
    """Assign one theme (and optionally one core subject) to many questions."""
    if not payload.f_ids:
        raise HTTPException(400, "no questions selected")
    if payload.theme_id is not None and db.get(Theme, payload.theme_id) is None:
        raise HTTPException(404, "theme not found")
    if payload.core_subject_id is not None:
        cs = db.get(CoreSubject, payload.core_subject_id)
        if cs is None:
            raise HTTPException(404, "core subject not found")
        if payload.theme_id is not None and cs.theme_id != payload.theme_id:
            raise HTTPException(409, "that core subject belongs to a different theme")

    rows = db.scalars(select(Fingerprint)
                      .where(Fingerprint.id.in_(payload.f_ids))).all()

    # Without a theme in the payload, each row keeps its own - and a core
    # subject only belongs under one theme. Applying it to rows themed
    # elsewhere would write the inconsistent pair that set_labels refuses to,
    # so those rows are skipped and counted rather than quietly corrupted.
    target_theme = (payload.theme_id if payload.theme_id is not None
                    else (cs.theme_id if payload.core_subject_id is not None else None))
    updated = wrong_theme = 0
    for f in rows:
        if payload.core_subject_id is not None and target_theme != (
                payload.theme_id if payload.theme_id is not None else f.theme_id):
            wrong_theme += 1
            continue
        if payload.theme_id is not None:
            # a core subject from the old theme cannot survive the move
            if f.theme_id != payload.theme_id and payload.core_subject_id is None:
                f.core_subject_id = None
            f.theme_id = payload.theme_id
        if payload.core_subject_id is not None:
            f.core_subject_id = payload.core_subject_id
        updated += 1
    db.commit()
    return {"updated": updated, "requested": len(payload.f_ids),
            "skipped_wrong_theme": wrong_theme}
